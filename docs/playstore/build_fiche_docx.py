#!/usr/bin/env python3
"""Fiche de publication Google Play de Xomiton au format Word (DOCX).

Lit docs/playstore/PUBLICATION_PLAY_STORE.md (source de vérité) et produit
docs/playstore/Fiche-publication-Play-Store-Xomiton.docx avec python-docx :
page de titre, sommaire (liens internes), titres, tableaux à en-tête coloré,
listes, citations, pied de page numéroté, puis l'image de présentation et les
miniatures des 8 captures avec leur texte alternatif en légende (et en texte
alternatif d'image pour les lecteurs d'écran).

Charte : titres Montserrat (repli automatique du logiciel si la police n'est pas
installée), texte Arial ; couleurs des jetons du design system
(design-system/xomiton/MASTER.md).

Usage : python3 docs/playstore/build_fiche_docx.py
Dépendance : pip install python-docx
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ICI = Path(__file__).resolve().parent
RACINE = ICI.parent.parent
SOURCE = ICI / "PUBLICATION_PLAY_STORE.md"
SORTIE = ICI / "Fiche-publication-Play-Store-Xomiton.docx"
LOGO = RACINE / "design-system" / "logo" / "png" / "logo-mono-blanc-1200w.png"

# Jetons de la charte (design-system/xomiton/MASTER.md)
PRIMAIRE = "0433C3"
DANGER = "E30E0E"
FOND = "FBF7FA"
TEXTE = "1B1630"
TEXTE_ATTENUE = "5B5670"
FILET = "E4DCE6"
SUCCES = "0B7A4B"
INFO_FOND = "EAF0FF"

POLICE_TITRES = "Montserrat"
POLICE_TEXTE = "Arial"
POLICE_CODE = "Consolas"


def rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


# --------------------------------------------------------------------------- bas niveau OOXML
def fond_cellule(cellule, hexa: str) -> None:
    tc_pr = cellule._tc.get_or_add_tcPr()
    for ancien in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(ancien)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa)
    tc_pr.append(shd)


def marges_cellule(cellule, haut=60, bas=60, gauche=90, droite=90) -> None:
    tc_pr = cellule._tc.get_or_add_tcPr()
    marges = OxmlElement("w:tcMar")
    for nom, val in (("top", haut), ("left", gauche), ("bottom", bas), ("right", droite)):
        el = OxmlElement(f"w:{nom}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        marges.append(el)
    tc_pr.append(marges)


def bordures_tableau(table, couleur=FILET, taille=4) -> None:
    tbl_pr = table._tbl.tblPr
    bordures = OxmlElement("w:tblBorders")
    for nom in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{nom}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(taille))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), couleur)
        bordures.append(el)
    tbl_pr.append(bordures)


def sans_bordures(table) -> None:
    tbl_pr = table._tbl.tblPr
    bordures = OxmlElement("w:tblBorders")
    for nom in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{nom}")
        el.set(qn("w:val"), "nil")
        bordures.append(el)
    tbl_pr.append(bordures)


def repeter_entete(ligne) -> None:
    tr_pr = ligne._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def ligne_indivisible(ligne) -> None:
    tr_pr = ligne._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def largeurs_colonnes(table, largeurs_cm: list[float]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    mise_en_page = OxmlElement("w:tblLayout")
    mise_en_page.set(qn("w:type"), "fixed")
    tbl_pr.append(mise_en_page)
    grille = table._tbl.tblGrid
    for i, col in enumerate(grille.findall(qn("w:gridCol"))):
        col.set(qn("w:w"), str(int(largeurs_cm[i] * 567)))
    for ligne in table.rows:
        for i, cellule in enumerate(ligne.cells):
            cellule.width = Cm(largeurs_cm[i])


def police(run, nom: str, taille: float | None = None, couleur: str | None = None,
           gras: bool | None = None, italique: bool | None = None) -> None:
    run.font.name = nom
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), nom)
    if taille is not None:
        run.font.size = Pt(taille)
    if couleur is not None:
        run.font.color.rgb = rgb(couleur)
    if gras is not None:
        run.font.bold = gras
    if italique is not None:
        run.font.italic = italique
    lang = run._r.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._r.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "fr-FR")


def champ(paragraphe, instruction: str, taille: float = 8.5, couleur: str = TEXTE_ATTENUE) -> None:
    """Champ Word (PAGE, NUMPAGES) mis à jour à l'ouverture."""
    run = paragraphe.add_run()
    police(run, POLICE_TEXTE, taille, couleur)
    debut = OxmlElement("w:fldChar")
    debut.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separation = OxmlElement("w:fldChar")
    separation.set(qn("w:fldCharType"), "separate")
    texte = OxmlElement("w:t")
    texte.text = "1"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for el in (debut, instr, separation, texte, fin):
        run._r.append(el)


def signet(paragraphe, nom: str, identifiant: int) -> None:
    debut = OxmlElement("w:bookmarkStart")
    debut.set(qn("w:id"), str(identifiant))
    debut.set(qn("w:name"), nom)
    fin = OxmlElement("w:bookmarkEnd")
    fin.set(qn("w:id"), str(identifiant))
    paragraphe._p.insert(0, debut)
    paragraphe._p.append(fin)


def lien_interne(paragraphe, texte: str, ancre: str, taille: float = 10.5, couleur: str = PRIMAIRE) -> None:
    lien = OxmlElement("w:hyperlink")
    lien.set(qn("w:anchor"), ancre)
    lien.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        r_fonts.set(qn(attr), POLICE_TEXTE)
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), couleur)
    r_pr.append(c)
    t = OxmlElement("w:sz")
    t.set(qn("w:val"), str(int(taille * 2)))
    r_pr.append(t)
    run.append(r_pr)
    el = OxmlElement("w:t")
    el.set(qn("xml:space"), "preserve")
    el.text = texte
    run.append(el)
    lien.append(run)
    paragraphe._p.append(lien)


def texte_alternatif(inline_shape, titre: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", titre)
    doc_pr.set("descr", description)


# --------------------------------------------------------------------------- mise en forme du texte
JETONS = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")
COMMENTAIRE = re.compile(r"<!--.*?-->")


def nettoyer(texte: str) -> str:
    texte = COMMENTAIRE.sub("", texte)
    texte = re.sub(r"<br\s*/?>", "\n", texte)
    return texte.strip()


LONG_JETON = re.compile(r"\S{20,}")


def couper(texte: str) -> str:
    """Autorise une coupe propre (espace de largeur nulle) après / . _ - dans les URL et codes longs."""
    def _coupe(m: re.Match) -> str:
        return re.sub(r"([/._\-?&=#])(?=[^/._\-?&=#])", "\\1\u200b", m.group(0))
    return LONG_JETON.sub(_coupe, texte)


def justifier(paragraphe, minimum: int = 110) -> None:
    """Justifie un paragraphe de corps assez long (jamais les titres, légendes ni cellules étroites)."""
    if len(paragraphe.text) >= minimum:
        paragraphe.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def langue_francaise(doc) -> None:
    """Langue fr-FR par défaut (styles) et césure automatique activée (réglages du document)."""
    styles = doc.styles.element
    defauts = styles.find(qn("w:docDefaults"))
    if defauts is not None:
        r_pr = defauts.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        for lang in r_pr.findall(qn("w:lang")):
            r_pr.remove(lang)
        lang = OxmlElement("w:lang")
        for attr in ("w:val", "w:eastAsia", "w:bidi"):
            lang.set(qn(attr), "fr-FR" if attr != "w:bidi" else "ar-SA")
        r_pr.append(lang)
    for style in doc.styles:
        r_pr = getattr(style.element, "rPr", None)
        if r_pr is not None and r_pr.find(qn("w:lang")) is None and style.type == 1:
            lang = OxmlElement("w:lang")
            lang.set(qn("w:val"), "fr-FR")
            r_pr.append(lang)
    reglages = doc.settings.element
    for nom in ("w:autoHyphenation", "w:doNotHyphenateCaps"):
        if reglages.find(qn(nom)) is None:
            el = OxmlElement(nom)
            reglages.append(el)


def ajouter_texte(paragraphe, texte: str, taille: float, couleur: str = TEXTE, gras: bool = False,
                  nom: str = POLICE_TEXTE) -> None:
    """Ajoute du texte avec **gras**, *italique* et `code` en ligne."""
    texte = couper(texte)
    for morceau in JETONS.split(texte):
        if not morceau:
            continue
        if morceau.startswith("**") and morceau.endswith("**") and len(morceau) > 4:
            run = paragraphe.add_run(morceau[2:-2])
            police(run, nom, taille, couleur, True)
        elif morceau.startswith("`") and morceau.endswith("`"):
            run = paragraphe.add_run(morceau[1:-1])
            police(run, POLICE_CODE, taille - 0.5, DANGER if False else PRIMAIRE, gras)
        elif morceau.startswith("*") and morceau.endswith("*") and len(morceau) > 2:
            run = paragraphe.add_run(morceau[1:-1])
            police(run, nom, taille, couleur, gras, True)
        else:
            run = paragraphe.add_run(morceau)
            police(run, nom, taille, couleur, gras)


def espacement(paragraphe, avant: float = 0, apres: float = 6, interligne: float = 1.25) -> None:
    fmt = paragraphe.paragraph_format
    fmt.space_before = Pt(avant)
    fmt.space_after = Pt(apres)
    fmt.line_spacing = interligne


# --------------------------------------------------------------------------- document
class Fiche:
    def __init__(self) -> None:
        self.doc = Document()
        self.signets = 0
        self.titres_h2: list[tuple[str, str]] = []  # (texte, signet)
        self.configurer()

    def configurer(self) -> None:
        section = self.doc.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)
        section.different_first_page_header_footer = True

        normal = self.doc.styles["Normal"]
        normal.font.name = POLICE_TEXTE
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE_TEXTE)

        for niveau, taille, couleur, avant, apres in ((1, 22, PRIMAIRE, 0, 10), (2, 16, PRIMAIRE, 20, 8),
                                                     (3, 12.5, TEXTE, 14, 5), (4, 11, TEXTE, 10, 4)):
            style = self.doc.styles[f"Heading {niveau}"]
            style.font.name = POLICE_TITRES
            style.font.size = Pt(taille)
            style.font.bold = True
            style.font.color.rgb = rgb(couleur)
            style.element.rPr.rFonts.set(qn("w:ascii"), POLICE_TITRES)
            style.element.rPr.rFonts.set(qn("w:hAnsi"), POLICE_TITRES)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE_TITRES)
            style.paragraph_format.space_before = Pt(avant)
            style.paragraph_format.space_after = Pt(apres)
            style.paragraph_format.keep_with_next = True

        # En-tête (pages suivantes) et pied de page numéroté.
        en_tete = section.header.paragraphs[0]
        en_tete.text = ""
        run = en_tete.add_run("Xomiton · Fiche de publication Google Play")
        police(run, POLICE_TITRES, 8.5, TEXTE_ATTENUE, True)
        pied = section.footer.paragraphs[0]
        pied.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = pied.add_run("Page ")
        police(run, POLICE_TEXTE, 8.5, TEXTE_ATTENUE)
        champ(pied, "PAGE")
        run = pied.add_run(" sur ")
        police(run, POLICE_TEXTE, 8.5, TEXTE_ATTENUE)
        champ(pied, "NUMPAGES")

    # ---- page de titre
    def page_de_titre(self, titre: str, version: str, date: str) -> None:
        bandeau = self.doc.add_table(rows=1, cols=1)
        bandeau.alignment = WD_TABLE_ALIGNMENT.CENTER
        sans_bordures(bandeau)
        largeurs_colonnes(bandeau, [17.0])
        cellule = bandeau.rows[0].cells[0]
        fond_cellule(cellule, PRIMAIRE)
        marges_cellule(cellule, 700, 700, 500, 500)
        p = cellule.paragraphs[0]
        if LOGO.exists():
            p.add_run().add_picture(str(LOGO), width=Cm(5.2))
        p.paragraph_format.space_after = Pt(60)

        p = cellule.add_paragraph()
        run = p.add_run("DOSSIER DE PUBLICATION")
        police(run, POLICE_TITRES, 11, "FFFFFF", True)
        espacement(p, 0, 6)
        p = cellule.add_paragraph()
        run = p.add_run("Fiche de publication Google Play")
        police(run, POLICE_TITRES, 28, "FFFFFF", True)
        espacement(p, 0, 8, 1.1)
        p = cellule.add_paragraph()
        run = p.add_run("Xomiton : écoute et sécurité")
        police(run, POLICE_TEXTE, 15, "FFFFFF")
        espacement(p, 0, 70)

        for etiquette, valeur in (("Version", version), ("Date", date), ("Application", "bj.xomiton.xomiton · 1.0.0 (code 1)"),
                                  ("Langue de la fiche", "Français uniquement")):
            p = cellule.add_paragraph()
            run = p.add_run(f"{etiquette}  ")
            police(run, POLICE_TITRES, 10, "C9D4FF", True)
            run = p.add_run(valeur)
            police(run, POLICE_TEXTE, 10.5, "FFFFFF")
            espacement(p, 0, 3)

        p = self.doc.add_paragraph()
        espacement(p, 14, 4)
        ajouter_texte(p, "Ce document est généré à partir de `PUBLICATION_PLAY_STORE.md`. Il ne remplace pas un avis juridique : "
                         "les pages légales et les déclarations de la console doivent être validées par le conseil du client avant le dépôt.",
                      9, TEXTE_ATTENUE)
        self.saut_de_page()

    def saut_de_page(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- sommaire
    def reserver_sommaire(self):
        titre = self.doc.add_paragraph(style="Heading 1")
        titre.add_run("Sommaire")
        # Les entrées sont écrites une fois tous les titres connus : paragraphe repère.
        repere = self.doc.add_paragraph()
        self.saut_de_page()
        return repere

    def remplir_sommaire(self, repere) -> None:
        suivant = repere
        for texte, ancre in self.titres_h2:
            nouveau = OxmlElement("w:p")
            suivant._p.addnext(nouveau)
            from docx.text.paragraph import Paragraph
            para = Paragraph(nouveau, repere._parent)
            espacement(para, 0, 5, 1.2)
            para.paragraph_format.left_indent = Cm(0.2)
            lien_interne(para, texte, ancre, 11)
            suivant = para
        repere._p.getparent().remove(repere._p)

    # ---- blocs
    def titre(self, niveau: int, texte: str) -> None:
        texte = nettoyer(texte)
        p = self.doc.add_paragraph(style=f"Heading {min(niveau, 4)}")
        ajouter_texte(p, texte, {1: 22, 2: 16, 3: 12.5}.get(niveau, 11), PRIMAIRE if niveau <= 2 else TEXTE, True, POLICE_TITRES)
        if niveau == 2:
            self.signets += 1
            nom = f"section{self.signets}"
            signet(p, nom, self.signets)
            self.titres_h2.append((re.sub(r"[*`]", "", texte), nom))
            # filet sous le titre
            p_pr = p._p.get_or_add_pPr()
            bordure = OxmlElement("w:pBdr")
            bas = OxmlElement("w:bottom")
            bas.set(qn("w:val"), "single")
            bas.set(qn("w:sz"), "8")
            bas.set(qn("w:space"), "4")
            bas.set(qn("w:color"), FILET)
            bordure.append(bas)
            p_pr.append(bordure)

    def paragraphe(self, texte: str) -> None:
        p = self.doc.add_paragraph()
        espacement(p, 0, 6)
        for i, ligne in enumerate(nettoyer(texte).split("\n")):
            if i:
                p.add_run().add_break()
            ajouter_texte(p, ligne, 10.5)
        justifier(p)

    def citation(self, texte: str) -> None:
        tableau = self.doc.add_table(rows=1, cols=1)
        sans_bordures(tableau)
        largeurs_colonnes(tableau, [17.0])
        cellule = tableau.rows[0].cells[0]
        fond_cellule(cellule, INFO_FOND)
        marges_cellule(cellule, 100, 100, 160, 140)
        tc_pr = cellule._tc.get_or_add_tcPr()
        bordures = OxmlElement("w:tcBorders")
        gauche = OxmlElement("w:left")
        gauche.set(qn("w:val"), "single")
        gauche.set(qn("w:sz"), "24")
        gauche.set(qn("w:color"), PRIMAIRE)
        bordures.append(gauche)
        tc_pr.append(bordures)
        p = cellule.paragraphs[0]
        espacement(p, 0, 0)
        ajouter_texte(p, nettoyer(texte), 10, TEXTE)
        justifier(p)
        espace = self.doc.add_paragraph()
        espacement(espace, 0, 2)

    def liste(self, elements: list[str], numerotee: bool) -> None:
        for i, texte in enumerate(elements, 1):
            p = self.doc.add_paragraph()
            espacement(p, 0, 3)
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            marque = f"{i}." if numerotee else "•"
            run = p.add_run(marque + "\t" if False else marque + " ")
            police(run, POLICE_TEXTE, 10.5, PRIMAIRE, True)
            ajouter_texte(p, nettoyer(texte), 10.5)

    def tableau(self, entete: list[str], lignes: list[list[str]]) -> None:
        colonnes = len(entete)
        taille = 9.5 if colonnes <= 3 else 8.5 if colonnes <= 5 else 7.5
        largeur_totale = 17.0
        # Largeurs proportionnelles à la longueur du contenu (bornées).
        poids = []
        for i in range(colonnes):
            longueur = max([len(nettoyer(entete[i]))] + [len(nettoyer(l[i])) for l in lignes if i < len(l)])
            poids.append(min(max(longueur, 8), 70) ** 0.8)
        total = sum(poids)
        minimum = 2.4 if colonnes >= 6 else 2.3 if colonnes >= 4 else 1.7
        largeurs = [max(minimum, largeur_totale * p / total) for p in poids]
        # Jamais plus étroit que le plus long mot de l'en-tête (évite « Dimension / s »).
        for i in range(colonnes):
            mot = max((len(w) for w in nettoyer(entete[i]).split()), default=4)
            largeurs[i] = max(largeurs[i], mot * 0.2 * (taille / 8.5) + 0.5)
        facteur = largeur_totale / sum(largeurs)
        largeurs = [l * facteur for l in largeurs]

        table = self.doc.add_table(rows=1, cols=colonnes)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordures_tableau(table)
        for i, texte in enumerate(entete):
            cellule = table.rows[0].cells[i]
            fond_cellule(cellule, PRIMAIRE)
            marges_cellule(cellule)
            p = cellule.paragraphs[0]
            espacement(p, 0, 0, 1.1)
            ajouter_texte(p, nettoyer(texte), taille - (0.5 if colonnes >= 6 else 0), "FFFFFF", True, POLICE_TITRES)
        repeter_entete(table.rows[0])
        for n, ligne in enumerate(lignes):
            cellules = table.add_row().cells
            ligne_indivisible(table.rows[-1])
            for i in range(colonnes):
                texte = ligne[i] if i < len(ligne) else ""
                fond_cellule(cellules[i], FOND if n % 2 else "FFFFFF")
                marges_cellule(cellules[i])
                p = cellules[i].paragraphs[0]
                espacement(p, 0, 0, 1.15)
                for k, partie in enumerate(nettoyer(texte).split("\n")):
                    if k:
                        p.add_run().add_break()
                    ajouter_texte(p, partie, taille, TEXTE, i == 0 and 2 < colonnes <= 5)
        largeurs_colonnes(table, largeurs)
        espace = self.doc.add_paragraph()
        espacement(espace, 0, 4)

    # ---- visuels
    def visuels(self, assets: dict) -> None:
        self.titre(3, "Aperçu des visuels")
        p = self.doc.add_paragraph()
        ajouter_texte(p, "Image de présentation (1024 × 500) et miniatures des 8 captures (1080 × 1920). "
                         "Le texte alternatif de chaque visuel est repris en légende.", 10, TEXTE_ATTENUE)
        espacement(p, 0, 6)

        presentation = assets["presentation"]
        if presentation["fichier"].exists():
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            forme = p.add_run().add_picture(str(presentation["fichier"]), width=Cm(13.5))
            texte_alternatif(forme, "Image de présentation", presentation["alt"])
            l = self.doc.add_paragraph()
            l.alignment = WD_ALIGN_PARAGRAPH.CENTER
            espacement(l, 2, 12)
            ajouter_texte(l, f"**Image de présentation.** {presentation['alt']}", 8.5, TEXTE_ATTENUE)

        captures = assets["captures"]
        par_ligne = 4
        table = self.doc.add_table(rows=0, cols=par_ligne)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sans_bordures(table)
        for debut in range(0, len(captures), par_ligne):
            ligne = table.add_row()
            ligne_indivisible(ligne)
            for i in range(par_ligne):
                cellule = ligne.cells[i]
                marges_cellule(cellule, 40, 80, 60, 60)
                if debut + i >= len(captures):
                    continue
                capture = captures[debut + i]
                p = cellule.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if capture["fichier"].exists():
                    forme = p.add_run().add_picture(str(capture["fichier"]), width=Cm(3.5))
                    texte_alternatif(forme, capture["titre"], capture["alt"])
                legende = cellule.add_paragraph()
                espacement(legende, 3, 0, 1.1)
                ajouter_texte(legende, f"**{capture['numero']}. {capture['titre']}**", 7.5, TEXTE)
                alt = cellule.add_paragraph()
                espacement(alt, 1, 0, 1.1)
                ajouter_texte(alt, capture["alt"], 7, TEXTE_ATTENUE)
        largeurs_colonnes(table, [17.0 / par_ligne] * par_ligne)
        espace = self.doc.add_paragraph()
        espacement(espace, 0, 6)


def charger_visuels() -> dict:
    config = json.loads((ICI / "visuels.json").read_text(encoding="utf-8"))
    presentation_alt = ""
    captures = []
    noms = []
    entrees = config.get("visuels") or config.get("captures") or []
    for entree in entrees:
        noms.append(entree)
    # Les textes alternatifs de référence sont ceux du tableau du Markdown (§10) : source de vérité de la fiche.
    markdown = SOURCE.read_text(encoding="utf-8")
    lignes = re.findall(r"^\| `assets/([^`]+)` \| [^|]+\| [^|]+\| ([^|]+)\| (.+) \|$", markdown, re.M)
    titres = dict(re.findall(r"(\d\d) « ([^»]+) »", re.search(r"Titres incrustés \(français\) : (.+)", markdown).group(1)))
    assets = {"presentation": {"fichier": ICI / "assets" / "feature-graphic-1024x500.png", "alt": ""}, "captures": []}
    for fichier, role, alt in lignes:
        alt = alt.strip()
        if fichier.startswith("feature-graphic"):
            assets["presentation"]["alt"] = alt
        elif fichier.startswith("screenshots/"):
            numero = fichier.split("/")[1][:2]
            assets["captures"].append({
                "fichier": ICI / "assets" / fichier,
                "numero": numero,
                "titre": titres.get(numero, fichier),
                "alt": alt,
            })
    return assets


def analyser(markdown: str):
    """Découpe le Markdown en blocs : (type, données)."""
    lignes = markdown.split("\n")
    i = 0
    blocs = []
    while i < len(lignes):
        ligne = lignes[i]
        if not ligne.strip():
            i += 1
            continue
        if ligne.startswith("#"):
            niveau = len(ligne) - len(ligne.lstrip("#"))
            blocs.append(("titre", (niveau, ligne[niveau:].strip())))
            i += 1
        elif ligne.startswith("|"):
            tableau = []
            while i < len(lignes) and lignes[i].startswith("|"):
                tableau.append(lignes[i])
                i += 1
            cellules = [[c.strip() for c in re.split(r"(?<!\\)\|", t.strip().strip("|"))] for t in tableau]
            cellules = [[c.replace("\\|", "|") for c in l] for l in cellules]
            if len(cellules) >= 2 and all(re.fullmatch(r":?-{3,}:?", c) for c in cellules[1]):
                blocs.append(("tableau", (cellules[0], cellules[2:])))
        elif ligne.startswith(">"):
            texte = []
            while i < len(lignes) and lignes[i].startswith(">"):
                texte.append(lignes[i].lstrip(">").strip())
                i += 1
            blocs.append(("citation", " ".join(texte)))
        elif re.match(r"^\s*(- |\* |• )", ligne) or re.match(r"^\s*\d+\. ", ligne):
            numerotee = bool(re.match(r"^\s*\d+\. ", ligne))
            elements = []
            marque = r"^\s*(- |\* |• |\d+\. )"
            while i < len(lignes) and (re.match(marque, lignes[i])
                                       or (lignes[i].startswith("  ") and lignes[i].strip() and elements)):
                texte = re.sub(marque, "", lignes[i])
                if lignes[i].startswith("  ") and not re.match(marque, lignes[i]):
                    elements[-1] += " " + lignes[i].strip()
                else:
                    elements.append(texte)
                i += 1
            blocs.append(("liste", (elements, numerotee)))
        else:
            texte = [ligne]
            i += 1
            while i < len(lignes) and lignes[i].strip() and not re.match(r"^(#|\||>|\s*- |\s*\* |\s*\d+\. |•)", lignes[i]):
                texte.append(lignes[i])
                i += 1
            blocs.append(("paragraphe", "\n".join(texte)))
    return blocs


def main() -> int:
    markdown = SOURCE.read_text(encoding="utf-8")
    entete = re.search(r"Version du dossier ([\d.]+) · (\d+ \w+ \d{4})", markdown)
    version, date = (entete.group(1), entete.group(2)) if entete else ("?", "?")
    blocs = analyser(markdown)
    assets = charger_visuels()

    fiche = Fiche()
    fiche.page_de_titre("Fiche de publication Google Play", version, date)
    repere = fiche.reserver_sommaire()

    visuels_places = False
    premier_titre = True
    for type_bloc, donnees in blocs:
        if type_bloc == "titre":
            niveau, texte = donnees
            if niveau == 1:
                continue  # le titre est sur la page de titre
            if niveau == 2 and not premier_titre:
                pass
            premier_titre = False
            fiche.titre(niveau, texte)
        elif type_bloc == "paragraphe":
            fiche.paragraphe(donnees)
        elif type_bloc == "citation":
            fiche.citation(donnees)
        elif type_bloc == "liste":
            fiche.liste(*donnees)
        elif type_bloc == "tableau":
            fiche.tableau(*donnees)
            # Visuels juste après le tableau des assets (§10).
            if not visuels_places and donnees[0] and donnees[0][0].lower().startswith("fichier") and any("feature-graphic" in l[0] for l in donnees[1]):
                fiche.visuels(assets)
                visuels_places = True

    if not visuels_places:
        print("Attention : tableau des assets introuvable, visuels non insérés.", file=sys.stderr)
        return 1

    fiche.remplir_sommaire(repere)
    fiche.doc.core_properties.title = "Fiche de publication Google Play – Xomiton"
    fiche.doc.core_properties.author = "Xomiton"
    fiche.doc.core_properties.subject = f"Version {version} du {date}"
    fiche.doc.core_properties.language = "fr-FR"
    langue_francaise(fiche.doc)
    fiche.doc.save(SORTIE)
    print(f"OK : {SORTIE.relative_to(RACINE)} (version {version}, {date}, {len(fiche.titres_h2)} sections, {len(assets['captures'])} captures)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
